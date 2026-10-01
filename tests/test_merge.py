"""merge_fdstools_mutect2_improved.py: both callers' tables -> one report."""
import argparse
import subprocess
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

import merge_fdstools_mutect2_improved as merge

AMPLICONS = {"mtNG_001": (19, 155), "mtNG_002": (133, 267), "mtNG_003": (259, 367),
             "mtNG_005": (470, 610), "mtNG_097": (16094, 16276)}
MT2_COLUMNS = ["MUTECT2", "vf_MT2", "rd_MT2", "MBQ", "Filter", "Pos", "Ref", "Variant", "GT", "Type"]


def fds(label, vf, depth, marker):
    """A row of the FDSTOOLS report table (vf in percent)."""
    start, end = AMPLICONS[marker]
    return {"FDSTOOLS": label, "vf_FDS": vf, "rd_FDS": round(vf * depth / 100),
            "interpolated_total_coverage": depth, "marker": marker,
            "marker_range": f"{start}-{end}", "num_markers": 1}


def fds_low(marker, reads):
    start, end = AMPLICONS[marker]
    return {"FDSTOOLS": "LOW", "rd_FDS": reads, "marker": marker, "marker_range": f"{start}-{end}"}


def mt2(label, vf, ad, pos, ref, alt, kind):
    """A row of the Mutect2 report table (vf as a fraction, ad as ref,alt reads)."""
    return dict(zip(MT2_COLUMNS, [label, vf, ad, "35,35", "PASS", pos, ref, alt, "0/1", kind]))


def write_inputs(tmp_path, fds_rows, mt2_rows, depths):
    library = tmp_path / "library.txt"
    library.write_text("[genome_position]\n" + "".join(
        f"{m} = chrM, {s}, {e}\n" for m, (s, e) in AMPLICONS.items()) + "[flanks]\n")
    fds_file, mt2_file, depth_file = tmp_path / "fds.txt", tmp_path / "mt2.txt", tmp_path / "depth.txt"
    pd.DataFrame(fds_rows).to_csv(fds_file, sep="\t", index=False)
    pd.DataFrame(mt2_rows, columns=MT2_COLUMNS).to_csv(mt2_file, sep="\t", index=False)
    # p09's depth at each amplicon's middle position; 500 reads unless given
    depth_file.write_text("".join(f"chrM\t{(s + e) // 2}\t{depths.get(m, 500)}\n"
                                  for m, (s, e) in AMPLICONS.items()))
    return fds_file, mt2_file, depth_file, library


def run(tmp_path, fds_rows, mt2_rows, depths={}, weighted=False, disabled=()):
    fds_file, mt2_file, depth_file, library = write_inputs(tmp_path, fds_rows, mt2_rows, depths)
    return merge.merge_variant_callers(str(fds_file), str(mt2_file), lh_thresh=10.0, min_vf=5.0,
                                       mutect2_depth_file=str(depth_file), marker_map=str(library),
                                       depth_threshold=10, weighted_average=weighted,
                                       mutect2_disabled=disabled)


def calls(report):
    return report[report["FMP"] != "LOW"].set_index("FMP")


@pytest.mark.parametrize("a, b", [("-309.1C", "-309.1c"), ("T16189C", "T16189Y"), ("T16189Y", "T16189H"),
                                  ("C756M", "C756Y"), ("A523-", "A523a")])
def test_merge_key_same_locus_any_spelling(a, b):
    assert merge.merge_key(a) == merge.merge_key(b)


def test_merge_key_keeps_substitution_and_deletion_apart():
    assert merge.merge_key("A523C") != merge.merge_key("A523-")


def test_agreeing_callers_share_a_row(tmp_path):
    report = calls(run(tmp_path, [fds("T152C", 100.0, 498, "mtNG_001")],
                       [mt2("T152C", 0.998, "1,497", 152, "T", "C", "SNP")]))
    assert len(report) == 1
    assert report.loc["T152C", "called_by_FDSTOOLS"] is True
    assert report.loc["T152C", "called_by_MUTECT2"] is True


def test_silent_caller_is_false(tmp_path):
    report = calls(run(tmp_path, [fds("A263G", 100.0, 900, "mtNG_003")], []))
    assert report.loc["A263G", "called_by_MUTECT2"] is False


def test_length_disagreement_decided_by_average(tmp_path):
    # FDSTOOLS 92% (major), Mutect2 80% (minor): average 86% < 90% -> minor
    report = calls(run(tmp_path, [fds("-309.1C", 92.0, 900, "mtNG_003")],
                       [mt2("-309.1c", 0.80, "180,720", 302, "A", "AC", "LHP")]))
    assert list(report.index) == ["-309.1c"]
    assert report.loc["-309.1c", "Type"] == "LHP"
    assert report.loc["-309.1c", "called_by_FDSTOOLS"] == "DISAGREEMENT"
    assert report.loc["-309.1c", "called_by_MUTECT2"] is True


@pytest.mark.parametrize("weighted, reported, kind, disagreeing", [
    (False, "T16189Y", "PHP", "called_by_FDSTOOLS"),   # rCRS T: (0 + 10) / 2 = 5%, present
    (True, "T16189C", "SNP", "called_by_MUTECT2"),     # 2000 FDSTOOLS reads outweigh 20 Mutect2 reads
])
def test_substitution_disagreement(tmp_path, weighted, reported, kind, disagreeing):
    # FDSTOOLS: no T; Mutect2: T on 2 of its 20 reads
    report = calls(run(tmp_path, [fds("T16189C", 100.0, 2000, "mtNG_097")],
                       [mt2("T16189Y", 0.90, "2,18", 16189, "T", "C", "PHP")], weighted=weighted))
    assert list(report.index) == [reported]
    assert report.loc[reported, "Type"] == kind
    assert report.loc[reported, disagreeing] == "DISAGREEMENT"


def fds_bases(label, vf, reads, depth, marker):
    """An FDSTOOLS row whose vf and reads are given per base ("C 40, A 10")."""
    return {**fds(label, 0.0, depth, marker), "vf_FDS": vf, "rd_FDS": reads}


def test_rows_per_base_become_one_row(tmp_path):
    # T 50%, C 40%, A 10% written as one row per base other than rCRS
    report = calls(run(tmp_path, [fds("T16189Y", 40.0, 500, "mtNG_097"), fds("T16189W", 10.0, 500, "mtNG_097")],
                       [mt2("T16189Y", 0.4, "250,200", 16189, "T", "C", "PHP"),
                        mt2("T16189W", 0.1, "250,50", 16189, "T", "A", "PHP")]))
    assert list(report.index) == ["T16189H"]
    row = report.loc["T16189H"]
    assert (row["vf_FDS"], row["rd_FDS"]) == ("C 40, A 10", "C 200, A 50")
    assert (row["vf_MT2"], row["rd_MT2"], row["Type"]) == ("C 0.4, A 0.1", "250,200,50", "PHP")
    assert row["called_by_FDSTOOLS"] is True and row["called_by_MUTECT2"] is True


def test_rows_per_base_without_rcrs(tmp_path):
    # C 70% and A 30%, no T: the old rows' codes both claim a T
    report = calls(run(tmp_path, [fds("T16189Y", 70.0, 500, "mtNG_097"), fds("T16189W", 30.0, 500, "mtNG_097")], []))
    assert list(report.index) == ["T16189M"]
    assert report.loc["T16189M", "vf_FDS"] == "C 70, A 30"


def test_a_base_one_caller_found_is_kept(tmp_path):
    # FDSTOOLS: T, C and A; Mutect2: T and C (A below its floor) -> H, Mutect2 disagrees
    report = calls(run(tmp_path, [fds_bases("T16189H", "C 40, A 10", "C 200, A 50", 500, "mtNG_097")],
                       [mt2("T16189Y", 0.44, "280,220", 16189, "T", "C", "PHP")]))
    assert list(report.index) == ["T16189H"]
    assert report.loc["T16189H", "called_by_FDSTOOLS"] is True
    assert report.loc["T16189H", "called_by_MUTECT2"] == "DISAGREEMENT"


def test_different_bases_from_the_two_callers_share_a_row(tmp_path):
    report = calls(run(tmp_path, [fds("T16189Y", 30.0, 500, "mtNG_097")],
                       [mt2("T16189W", 0.08, "460,40", 16189, "T", "A", "PHP")]))
    assert list(report.index) == ["T16189H"]
    assert report.loc["T16189H", "called_by_FDSTOOLS"] == "DISAGREEMENT"
    assert report.loc["T16189H", "called_by_MUTECT2"] == "DISAGREEMENT"


@pytest.mark.parametrize("deleted, reported", [(None, "T16189Y"), (10.0, "T16189C")])
def test_rcrs_share_leaves_out_deleted_molecules(tmp_path, deleted, reported):
    # FDSTOOLS: C 90%, the remaining 10% T or deleted; Mutect2: C 93%, T 7%
    fds_rows = [fds("T16189C", 90.0, 500, "mtNG_097")]
    if deleted:
        fds_rows.append(fds("T16189t", deleted, 500, "mtNG_097"))
    report = calls(run(tmp_path, fds_rows, [mt2("T16189Y", 0.93, "35,465", 16189, "T", "C", "PHP")]))
    assert reported in report.index


def test_mutect2_rcrs_share_from_its_reads(tmp_path):
    # 24 reads, all G: Mutect2 reports 0.96, which would leave 4% for rCRS and average
    # 2% with FDSTOOLS' 0%; its reads show no rCRS
    fds_file, mt2_file, depth_file, library = write_inputs(
        tmp_path, [fds("A73G", 100.0, 100, "mtNG_001")], [mt2("A73R", 0.96, "0,24", 73, "A", "G", "PHP")], {})
    report = calls(merge.merge_variant_callers(str(fds_file), str(mt2_file), lh_thresh=10.0, min_vf=2.0,
                                               mutect2_depth_file=str(depth_file), marker_map=str(library)))
    assert list(report.index) == ["A73G"]
    assert report.loc["A73G", "called_by_MUTECT2"] == "DISAGREEMENT"


def test_rcrs_missing_with_two_other_bases(tmp_path):
    report = calls(run(tmp_path, [fds_bases("T16189M", "C 60, A 40", "C 300, A 200", 500, "mtNG_097")],
                       [mt2("T16189C", 0.96, "20,480", 16189, "T", "C", "SNP")]))
    assert list(report.index) == ["T16189M"]
    assert report.loc["T16189M", "Type"] == "PHP"
    assert report.loc["T16189M", "called_by_MUTECT2"] == "DISAGREEMENT"


def test_values_per_base_do_not_stop_the_other_rows(tmp_path):
    # one row with vf per base makes the column text; the length disagreement on
    # another row is still averaged (it crashed on the old branch at min_vf 1)
    report = calls(run(tmp_path, [fds_bases("A16265V", "C 6, G 4", "C 30, G 20", 500, "mtNG_097"),
                                  fds("-309.1C", 92.0, 900, "mtNG_003")],
                       [mt2("-309.1c", 0.80, "180,720", 302, "A", "AC", "LHP")]))
    assert "-309.1c" in report.index
    assert report.loc["A16265V", "vf_FDS"] == "C 6, G 4"


def test_minor_deletion_from_both_callers_shares_a_row(tmp_path):
    report = calls(run(tmp_path, [fds("A523a", 40.0, 500, "mtNG_005")],
                       [mt2("A523a", 0.40, "300,200", 513, "GCA", "G", "LHP")]))
    assert list(report.index) == ["A523a"]
    assert report.loc["A523a", "called_by_MUTECT2"] is True


def test_deletion_major_minor_split_is_a_disagreement(tmp_path):
    # FDSTOOLS 95% (A523-), Mutect2 85% (A523a): average 90% -> major
    report = calls(run(tmp_path, [fds("A523-", 95.0, 500, "mtNG_005")],
                       [mt2("A523a", 0.85, "75,425", 513, "GCA", "G", "LHP")]))
    assert list(report.index) == ["A523-"]
    assert report.loc["A523-", "Type"] == "DEL"
    assert report.loc["A523-", "called_by_MUTECT2"] == "DISAGREEMENT"


def test_low_rows_per_amplicon_for_both_callers(tmp_path):
    report = run(tmp_path, [fds("T152C", 100.0, 498, "mtNG_001"), fds_low("mtNG_002", 6)], [],
                 depths={"mtNG_003": 3})
    low = report[report["FMP"] == "LOW"].set_index("marker")
    assert sorted(low.index) == ["mtNG_002", "mtNG_003"]
    assert (low.loc["mtNG_002", "called_by_FDSTOOLS"], low.loc["mtNG_002", "called_by_MUTECT2"]) == ("LOW", "OK")
    assert (low.loc["mtNG_003", "called_by_FDSTOOLS"], low.loc["mtNG_003", "called_by_MUTECT2"]) == ("OK", "LOW")
    assert low.loc["mtNG_003", "rd_MT2"] == 3


def test_script_writes_excel_with_flags_coloured(tmp_path):
    fds_file, mt2_file, depth_file, library = write_inputs(
        tmp_path,
        [fds("A263G", 100.0, 900, "mtNG_003"), fds("-309.1C", 92.0, 900, "mtNG_003"), fds_low("mtNG_002", 6)],
        [mt2("-309.1c", 0.80, "180,720", 302, "A", "AC", "LHP")], {})
    xlsx = tmp_path / "merged.xlsx"
    subprocess.run([sys.executable, merge.__file__, str(fds_file), str(mt2_file), str(xlsx),
                    "--mutect2_depth", str(depth_file), "--marker_map", str(library)], check=True)
    ws = load_workbook(xlsx).active
    header = [c.value for c in ws[1]]
    col = {name: header.index(name) for name in ("FMP", "called_by_FDSTOOLS", "called_by_MUTECT2")}
    rows = {r[col["FMP"]].value: r for r in ws.iter_rows(min_row=2)}
    fill = lambda cell: cell.fill.start_color.rgb[-6:]
    assert [r[col["FMP"]].value for r in ws.iter_rows(min_row=2)][-1] == "LOW"
    assert fill(rows["A263G"][col["called_by_MUTECT2"]]) == "FFC7CE"          # missed call: red
    assert fill(rows["-309.1c"][col["called_by_FDSTOOLS"]]) == "FFD966"       # disagreement: its own colour
    assert rows["LOW"][col["called_by_FDSTOOLS"]].value == "LOW"
    assert fill(rows["LOW"][col["called_by_FDSTOOLS"]]) == "FFC7CE"          # low coverage: red, like a miss
    assert fill(rows["LOW"][col["FMP"]]) == "FEFE01"                          # the LOW row itself stays yellow
    assert rows["LOW"][col["called_by_MUTECT2"]].value == "OK"
    assert fill(rows["LOW"][col["called_by_MUTECT2"]]) != "FFC7CE"           # ok is not a miss


def test_script_writes_values_per_base_in_percent(tmp_path):
    fds_file, mt2_file, depth_file, library = write_inputs(
        tmp_path, [fds_bases("T16189H", "C 40, A 10", "C 200, A 50", 500, "mtNG_097")],
        [mt2("T16189H", "C 0.4, A 0.1", "250,200,50", 16189, "T", "C", "PHP")], {})
    xlsx = tmp_path / "merged.xlsx"
    subprocess.run([sys.executable, merge.__file__, str(fds_file), str(mt2_file), str(xlsx),
                    "--mutect2_depth", str(depth_file), "--marker_map", str(library)], check=True)
    ws = load_workbook(xlsx).active
    header = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "T16189H")
    assert row[header.index("vf_MT2")].value == "C 40, A 10"
    assert row[0].fill.start_color.rgb[-6:] == "FAC000"                        # an IUPAC code


# 16180-16193 with T16189C: FDSTOOLS sees one C more on 30% of the molecules; Mutect2
# adds a boundary PHP of its own, T16195Y lies just past the region and A263G far outside
C_STRETCH_FDS = [fds("T16189C", 100.0, 900, "mtNG_097"), fds("-16193.1c", 30.0, 900, "mtNG_097"),
                 fds("A263G", 100.0, 900, "mtNG_003")]
C_STRETCH_MT2 = [mt2("T16189C", 0.99, "5,495", 16189, "T", "C", "SNP"),
                 mt2("A16183M", 0.07, "465,35", 16183, "A", "C", "PHP"),
                 mt2("-16193.1c", 0.34, "330,170", 16193, "C", "CC", "LHP"),
                 mt2("T16195Y", 0.20, "400,100", 16195, "T", "C", "PHP")]


def test_dominant_molecule_passes_through(tmp_path):
    report = calls(run(tmp_path, [{**fds("T16189C", 100.0, 900, "mtNG_097"), "dominant_molecule": "T16189C (62.0%)"}],
                       [mt2("T16189C", 0.99, "5,495", 16189, "T", "C", "SNP")]))
    assert report.loc["T16189C", "dominant_molecule"] == "T16189C (62.0%)"


def test_mutect2_calls_every_region_by_default(tmp_path):
    report = calls(run(tmp_path, C_STRETCH_FDS, C_STRETCH_MT2))
    assert report.loc["A16183M", "called_by_FDSTOOLS"] is False
    assert report.loc["-16193.1c", "called_by_MUTECT2"] is True


def test_disabled_region_keeps_mutect2_majors_and_notes_its_minor_calls(tmp_path):
    report = calls(run(tmp_path, C_STRETCH_FDS, C_STRETCH_MT2, disabled=["16180-16193"]))
    assert "A16183M" not in report.index
    assert report.loc["T16189C", "called_by_MUTECT2"] is True
    assert report.loc["-16193.1c", "called_by_MUTECT2"] == "DISABLED"
    assert pd.isna(report.loc["-16193.1c", "MUTECT2"])
    assert report.loc["A263G", "called_by_MUTECT2"] is False
    assert report.loc["T16195Y", "called_by_FDSTOOLS"] is False
    assert pd.isna(report.loc["T16195Y", "variant_note"])
    note = "Mutect2 minor calls left out: A16183M 7.0%, -16193.1c 34.0%"
    assert report.loc["T16189C", "variant_note"] == note
    assert report.loc["-16193.1c", "variant_note"] == note
    assert pd.isna(report.loc["A263G", "variant_note"])


@pytest.mark.parametrize("value, names", [("none", []), ("all", ["16180-16193", "300-315", "57-60"]),
                                          ("300-315, 16180-16193", ["300-315", "16180-16193"])])
def test_disabled_regions_option(value, names):
    assert merge.disabled_regions(value) == names


def test_disabled_regions_option_rejects_unknown_region():
    with pytest.raises(argparse.ArgumentTypeError):
        merge.disabled_regions("452-463")


def test_script_disables_mutect2_in_a_region(tmp_path):
    fds_file, mt2_file, depth_file, library = write_inputs(tmp_path, C_STRETCH_FDS, C_STRETCH_MT2, {})
    xlsx = tmp_path / "merged.xlsx"
    subprocess.run([sys.executable, merge.__file__, str(fds_file), str(mt2_file), str(xlsx),
                    "--mutect2_depth", str(depth_file), "--marker_map", str(library),
                    "--mutect2_disabled_regions", "16180-16193"], check=True)
    ws = load_workbook(xlsx).active
    header = [c.value for c in ws[1]]
    col = {name: header.index(name) for name in ("FMP", "called_by_MUTECT2")}
    rows = {r[col["FMP"]].value: r for r in ws.iter_rows(min_row=2)}
    fill = lambda cell: cell.fill.start_color.rgb[-6:]
    assert rows["-16193.1c"][col["called_by_MUTECT2"]].value == "DISABLED"
    assert fill(rows["-16193.1c"][col["called_by_MUTECT2"]]) != "FFC7CE"   # disabled is not a miss
    assert fill(rows["-16193.1c"][col["FMP"]]) != "F50003"
    assert fill(rows["A263G"][col["called_by_MUTECT2"]]) == "FFC7CE"
